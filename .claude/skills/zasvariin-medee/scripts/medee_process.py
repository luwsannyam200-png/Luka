#!/usr/bin/env python3
"""Засварын мэдээ: clean the Mine2TL repair log the way it is done by hand.

Reproduces the manual workbook ("ЗМ 27.xlsx"):
  1  drop column "Хэсгийн №"
  2  Дуудлага: location = the mechanic's location, service = "Дуудлага"
  3  drop the trailing total row
  4  Төлөвлөгөөт: service value -> "Дис тайлбар", service = "Төлөвлөгөөт"
  5  no driver: service = "Ачигч"
  5a Төлөвлөгөөт + Ачигч + empty "Дис тайлбар": ТҮ-x from "Мех тайлбар" is only
     suggested; it is written only if confirmed with --tu
  6  location "Замд": location = the mechanic's location
  7  checks: services include Ачигч/Дуудлага/Төлөвлөгөөт, none empty, only УХ/ЦХ
  8  red rows count as blue; blue (unfinished) rows first; their end = day of the
     first finished row at 06:00
  9-14 hours: "Дууссан огноо" = spent, "Засварын код" = waiting,
     "Машины төлөв" = total (see README in SKILL.md)

Usage:
  medee_process.py --input "засварын мэдээ.xlsx" --out "ZM_27.xlsx"
      [--mech mech.json] [--tu tu.json] [--end "2026-09-27 06:00"] [--check-only]
mech.json: {"Tsedenjav.I": "Ухаа худаг", ...} overrides the detected locations.
tu.json:   {"6494MMA": "ТҮ-2"} confirmed 5a values.
--check-only prints the detected mechanics / suggestions without writing.
"""
import argparse
import copy
import datetime as dt
import json
import re
import sys
import zipfile
from collections import Counter, defaultdict

import openpyxl
from openpyxl.styles import PatternFill
from openpyxl.worksheet.table import Table, TableStyleInfo

UX, CX = "Ухаа худаг", "Цагаан хад"
LOCS = (UX, CX)
BLUE = "FF2EAAFC"
GRAY = "FFD6D6D6"
ACC = '_(* #,##0.00_);_(* \\(#,##0.00\\);_(* "-"??_);_(@_)'
WIDTHS = [15.1, 12.6, 17.4, 8.6, 12.8, 17.1, 10.6, 12.1, 12.3, 25.3, 15.9, 9.2, 12.7, 14.9, 15.4, 12.6, 17.2, 13.4, 13.8, 13.4, 14.0]
TU_RE = re.compile(r"т[үуy]\s*-?\s*(\d(?:\.\d)?)", re.I)
EPOCH = dt.datetime(1899, 12, 30)


def norm(s):
    return re.sub(r"\s+", " ", str(s or "")).strip().lower()


def serial(d):
    return (d - EPOCH).total_seconds() / 86400


def hm_hours(t):
    """Excel's coercion of "hh:mm" text times 24 (blank -> 0)."""
    if t in (None, ""):
        return 0.0
    m = re.fullmatch(r"\s*(\d+):(\d{1,2})(?::(\d{1,2}))?\s*", str(t))
    if not m:
        raise ValueError(f"цагийн утга танигдсангүй: {t!r}")
    return int(m.group(1)) + int(m.group(2)) / 60 + int(m.group(3) or 0) / 3600


def is_reddish(rgb):
    if not rgb or len(rgb) < 6:
        return False
    r, g, b = int(rgb[-6:-4], 16), int(rgb[-4:-2], 16), int(rgb[-2:], 16)
    return r > 180 and g < 120 and b < 120


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--input", required=True)
    ap.add_argument("--out")
    ap.add_argument("--mech")
    ap.add_argument("--tu")
    ap.add_argument("--end")
    ap.add_argument("--check-only", action="store_true")
    a = ap.parse_args()

    ws = openpyxl.load_workbook(a.input).worksheets[0]
    header = [c.value for c in ws[1]]
    names = [norm(h) for h in header]
    need = ["засварын төрөл", "техникийн №", "жолооч/нэр, овог/", "засварын байршил", "үйлчилгээ", "дис тайлбар",
            "мех тайлбар", "механик", "хүлээсэн цаг", "зарцуулсан цаг", "хүсэлт өгсөн огноо", "эхэлсэн огноо",
            "дууссан огноо", "засварын код", "машины төлөв"]
    missing = [n for n in need if n not in names]
    if missing:
        sys.exit("Багана олдсонгүй: " + ", ".join(missing))
    drop = names.index("хэсгийн №") if "хэсгийн №" in names else None  # step 1
    keep = [i for i in range(len(header)) if i != drop]
    H = [header[i] for i in keep]
    I = {norm(h): j for j, h in enumerate(H)}

    # read rows (values + source cell styles); step 3 drops rows without a repair type
    rows, dropped = [], 0
    for r in ws.iter_rows(min_row=2):
        vals = [r[i].value for i in keep]
        if not vals[I["засварын төрөл"]] or not vals[I["техникийн №"]]:
            if any(c.value is not None for c in r):
                dropped += 1
            continue
        rows.append({"v": vals, "cells": [r[i] for i in keep], "fill": r[0].fill.fgColor.rgb if r[0].fill.fill_type else None})

    g = lambda row, k: row["v"][I[k]]

    def s(row, k, v):
        row["v"][I[k]] = v

    # state per row
    warnings = []
    for row in rows:
        end = g(row, "дууссан огноо")
        row["red"] = is_reddish(row["fill"])
        row["blue"] = end is None
        if row["fill"] == GRAY and end is None:
            warnings.append(f'{g(row, "техникийн №")}: саарал өнгөтэй боловч Дууссан огноо хоосон, дуусаагүй гэж тооцов')
        if (row["fill"] == BLUE or row["red"]) and end is not None:
            warnings.append(f'{g(row, "техникийн №")}: цэнхэр/улаан боловч Дууссан огноотой, дууссан гэж тооцов')

    # mechanic locations: auto from non-Замд work, overridable
    seen = defaultdict(Counter)
    for row in rows:
        if g(row, "засварын байршил") in LOCS:
            seen[g(row, "механик")][g(row, "засварын байршил")] += 1
    mech = {m: c.most_common(1)[0][0] for m, c in seen.items() if m}
    overrides = json.load(open(a.mech, encoding="utf-8")) if a.mech else {}
    mech.update(overrides)
    needed = sorted({g(r, "механик") for r in rows if g(r, "засварын төрөл") == "Дуудлага" or g(r, "засварын байршил") == "Замд"}, key=str)
    unknown = [m for m in needed if m not in mech]

    # 5a suggestions
    tu_conf = json.load(open(a.tu, encoding="utf-8")) if a.tu else {}
    suggestions = []
    for row in rows:
        if g(row, "засварын төрөл") == "Төлөвлөгөөт" and not g(row, "жолооч/нэр, овог/") and not g(row, "үйлчилгээ") and not g(row, "дис тайлбар"):
            m = TU_RE.search(g(row, "мех тайлбар") or "")
            suggestions.append({"tech": g(row, "техникийн №"), "mex": g(row, "мех тайлбар"), "suggest": f"ТҮ-{m.group(1)}" if m else None})

    report = {
        "rows": len(rows), "dropped_rows": dropped,
        "mechanics": {m: {"location": mech.get(m), "auto": seen[m].most_common(1)[0][0] if seen.get(m) else None,
                          "override": m in overrides, "counts": dict(seen.get(m, {}))} for m in sorted(set(mech) | set(needed), key=str)},
        "mechanics_needed": needed, "mechanics_unknown": unknown,
        "tu_suggestions": suggestions, "red_rows": sum(r["red"] for r in rows), "warnings": warnings,
    }
    if a.check_only or unknown or not a.out:
        if unknown:
            report["error"] = "Байршил тодорхойгүй механик: " + ", ".join(map(str, unknown)) + " (--mech ашиглана уу)"
        json.dump(report, sys.stdout, ensure_ascii=False, indent=1, default=str)
        print()
        sys.exit(1 if unknown else 0)

    # steps 2, 4, 5, 5a, 6 in the manual order
    for row in rows:
        if g(row, "засварын төрөл") == "Дуудлага":
            s(row, "засварын байршил", mech[g(row, "механик")])
            s(row, "үйлчилгээ", "Дуудлага")
    for row in rows:
        if g(row, "засварын төрөл") == "Төлөвлөгөөт":
            s(row, "дис тайлбар", g(row, "үйлчилгээ"))
            s(row, "үйлчилгээ", "Төлөвлөгөөт")
    for row in rows:
        if not g(row, "жолооч/нэр, овог/"):
            s(row, "үйлчилгээ", "Ачигч")
    for row in rows:
        t = tu_conf.get(g(row, "техникийн №"))
        if t and g(row, "засварын төрөл") == "Төлөвлөгөөт" and not g(row, "дис тайлбар"):
            s(row, "дис тайлбар", t)
    for row in rows:
        if g(row, "засварын байршил") == "Замд":
            s(row, "засварын байршил", mech[g(row, "механик")])

    # step 7 checks
    services = Counter(g(r, "үйлчилгээ") for r in rows)
    locs = Counter(g(r, "засварын байршил") for r in rows)
    problems = []
    for want in ("Ачигч", "Дуудлага", "Төлөвлөгөөт"):
        if not services.get(want):
            warnings.append(f'Үйлчилгээнд "{want}" алга')
    if services.get(None) or services.get(""):
        problems.append("Үйлчилгээ хоосон мөр үлдсэн")
    bad_locs = [l for l in locs if l not in LOCS]
    if bad_locs:
        problems.append(f"Байршилд УХ/ЦХ-ээс өөр утга: {bad_locs}")

    # step 8: order and end time
    blue = [r for r in rows if r["blue"]]
    gray = [r for r in rows if not r["blue"]]
    if a.end:
        end_dt = dt.datetime.fromisoformat(a.end)
    elif gray:
        end_dt = dt.datetime.combine(g(gray[0], "дууссан огноо").date(), dt.time(6, 0))
    else:
        sys.exit("Дууссан засвар (саарал мөр) алга, --end өгнө үү")
    for r in blue:
        if g(r, "эхэлсэн огноо") and g(r, "эхэлсэн огноо") > end_dt:
            warnings.append(f'{g(r, "техникийн №")}: {end_dt:%H:%M}-аас хойш эхэлсэн')
    ordered = blue + gray

    # steps 9-14 (formulas as in the manual workbook, with cached values)
    col = lambda k: I[k] + 1
    ref = lambda k: f"Table1[[#This Row],[{H[I[k]]}]]"
    cached = {}
    for r in blue:
        st, rq = g(r, "эхэлсэн огноо"), g(r, "хүсэлт өгсөн огноо")
        if st is None or rq is None:
            problems.append(f'{g(r, "техникийн №")}: Эхэлсэн эсвэл Хүсэлт өгсөн огноо хоосон')
            continue
        wait = (serial(st) - serial(rq)) * 24
        spent = (serial(end_dt) - serial(st)) * 24
        r["calc"] = {"механик": ("=" + ref("эхэлсэн огноо") + "*24-" + ref("хүсэлт өгсөн огноо") + "*24", wait),
                     "хүлээсэн цаг": ("=" + ref("дууссан огноо") + "*24-" + ref("эхэлсэн огноо") + "*24", (spent - serial(st)) * 24),
                     "зарцуулсан цаг": (None, wait), "дууссан огноо": (None, spent), "засварын код": (None, wait),
                     "машины төлөв": ("=" + ref("дууссан огноо") + "+" + ref("засварын код"), spent + wait)}
    for r in gray:
        sp, wt = hm_hours(g(r, "зарцуулсан цаг")), hm_hours(g(r, "хүлээсэн цаг"))
        r["calc"] = {"дууссан огноо": ("=" + ref("зарцуулсан цаг") + "*24", sp),
                     "засварын код": ("=" + ref("хүлээсэн цаг") + "*24", wt),
                     "машины төлөв": ("=" + ref("дууссан огноо") + "+" + ref("засварын код"), sp + wt)}
    if problems:
        report["error"] = problems
        json.dump(report, sys.stdout, ensure_ascii=False, indent=1, default=str)
        print()
        sys.exit(1)

    wb = openpyxl.Workbook()
    out = wb.active
    out.title = "Sheet"
    for j, h in enumerate(H, 1):
        c = out.cell(1, j, h)
        src = ws.cell(1, keep[j - 1] + 1)
        c.font, c.fill, c.border, c.alignment = copy.copy(src.font), copy.copy(src.fill), copy.copy(src.border), copy.copy(src.alignment)
        c.number_format = src.number_format
        out.column_dimensions[openpyxl.utils.get_column_letter(j)].width = WIDTHS[j - 1] if j <= len(WIDTHS) else 12
    blue_fill = PatternFill("solid", fgColor=BLUE)
    for i, r in enumerate(ordered, 2):
        for j in range(len(H)):
            c = out.cell(i, j + 1)
            src = r["cells"][j]
            c.font, c.border, c.alignment = copy.copy(src.font), copy.copy(src.border), copy.copy(src.alignment)
            c.fill = blue_fill if r["blue"] else copy.copy(src.fill)
            c.number_format = src.number_format
            c.value = r["v"][j]
        for k, (formula, val) in r["calc"].items():
            c = out.cell(i, col(k))
            if formula:
                c.value = formula
                cached[c.coordinate] = val
            else:
                c.value = val
            c.number_format = ACC if k in ("дууссан огноо", "засварын код", "машины төлөв") else "General"
    last = openpyxl.utils.get_column_letter(len(H))
    t = Table(displayName="Table1", ref=f"A1:{last}{len(ordered) + 1}")
    t.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
    out.add_table(t)
    wb.save(a.out)
    add_cached_values(a.out, cached)

    report.update({"end": str(end_dt), "blue": len(blue), "gray": len(gray), "services": services, "locations": locs,
                   "warnings": warnings})
    json.dump(report, sys.stdout, ensure_ascii=False, indent=1, default=str)
    print()


def add_cached_values(path, cached):
    """openpyxl writes formulas without results; add them so every viewer shows numbers."""
    z = zipfile.ZipFile(path)
    files = {n: z.read(n) for n in z.namelist()}
    infos = {n: z.getinfo(n) for n in z.namelist()}
    p = "xl/worksheets/sheet1.xml"
    xml = files[p].decode()

    def fill(m):
        ref = m.group(1)
        if ref not in cached:
            return m.group(0)
        return f'<c r="{ref}"{m.group(2)}><f>{m.group(3)}</f><v>{repr(float(cached[ref]))}</v></c>'
    xml = re.sub(r'<c r="([A-Z]+\d+)"([^>]*)><f>(.*?)</f><v\s*/?>(?:</v>)?</c>', fill, xml)
    xml = re.sub(r'<c r="([A-Z]+\d+)"([^>]*)><f>(.*?)</f></c>', fill, xml)
    files[p] = xml.encode()
    with zipfile.ZipFile(path + ".tmp", "w", zipfile.ZIP_DEFLATED) as o:
        for n in z.namelist():
            o.writestr(infos[n], files[n])
    import os
    os.replace(path + ".tmp", path)


if __name__ == "__main__":
    main()
