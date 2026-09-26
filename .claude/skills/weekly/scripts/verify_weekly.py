#!/usr/bin/env python3
"""Check a workbook produced by build_weekly.py against its original.

  verify_weekly.py --original "Weekly draft.xlsx" --new "Weekly draft_09.26.xlsx" --date 09.26

Fails (exit 1) if any existing sheet changed (values or formulas), the zip is
broken, or the new sheet's totals do not equal the sum of their parts.
"""
import argparse
import sys
import zipfile

import openpyxl


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--original", required=True)
    ap.add_argument("--new", required=True)
    ap.add_argument("--date", required=True)
    a = ap.parse_args()
    problems = []

    if zipfile.ZipFile(a.new).testzip() is not None:
        problems.append("zip is corrupt")

    for data_only in (True, False):
        o = openpyxl.load_workbook(a.original, data_only=data_only)
        n = openpyxl.load_workbook(a.new, data_only=data_only)
        if n.sheetnames[:-1] != o.sheetnames or n.sheetnames[-1] != a.date:
            problems.append(f"sheet list unexpected: {n.sheetnames[-3:]}")
        for s in o.sheetnames:
            if s not in n.sheetnames:
                problems.append(f"sheet {s} missing")
                continue
            for i, (x, y) in enumerate(zip(o[s].iter_rows(values_only=True), n[s].iter_rows(values_only=True)), 1):
                if x != y:
                    problems.append(f"{s} row {i} changed ({'values' if data_only else 'formulas'})")
                    break

    n = openpyxl.load_workbook(a.new, data_only=True)
    t = n[a.date]
    hdr = {}
    for c in range(1, t.max_column + 1):
        for r in (1, 2):
            v = str(t.cell(r, c).value or "").strip().lower()
            if v:
                hdr.setdefault(v, c)
    ux, cx, zbn = hdr["ух"], hdr["цх"], hdr["збн"]
    sup = [hdr["aode"], hdr["очлуур од"], hdr["parts and oil"]]
    total, allw, tlw = hdr["нийт агуулах"], hdr["all warehouse"], hdr["tl warehouse"]
    for r in range(3, t.max_row + 1):
        if t.cell(r, hdr["item code"]).value is None:
            continue
        parts = [t.cell(r, c).value or 0 for c in [ux, cx] + sup]
        e = t.cell(r, zbn).value
        checks = [(total, sum(parts), "Нийт агуулах")]
        if e:
            checks += [(allw, sum(parts) / e, "All warehouse"), (tlw, (parts[0] + parts[1]) / e, "TL Warehouse")]
        for c, expected, label in checks:
            got = t.cell(r, c).value
            if not isinstance(got, (int, float)) or abs(got - expected) > 1e-4:
                problems.append(f"row {r}: {label} = {got!r}, expected {expected}")

    if problems:
        print("FAIL")
        print("\n".join(problems[:50]))
        sys.exit(1)
    print(f"OK: {len(o.sheetnames)} old sheets unchanged, sheet {a.date} consistent")


if __name__ == "__main__":
    main()
