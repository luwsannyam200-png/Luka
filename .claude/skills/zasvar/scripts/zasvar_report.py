#!/usr/bin/env python3
"""Repair request report from the Mine2TL export ("засвар ...rqst.xlsx").

Output workbook:
  Pivot     rows Засварын төрөл > Засварын байршил > Үйлчилгээ, columns = request day, values = count
  Хүлээлт   same rows, columns = days waiting (0, 1-3, 4-7, 8+)
  Анхаарах  urgent requests, long waits, ERP WO numbers, data errors

Usage:
  zasvar_report.py --input "засвар 9-26rqst.xlsx" --out "Засварын хүсэлт_09.27.xlsx" [--today 2026-09-27]
--today defaults to the day of the latest request in the file.
"""
import argparse
import datetime as dt
import json
import re
import sys
from collections import Counter, defaultdict

import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

BLANK = "(хоосон)"
URGENT_WORDS = re.compile(r"тормоз|тоормоз|tormoz|brake|гал\s*авсан|утаа", re.I)
WO_RE = re.compile(r"\bER\d{6,8}\b", re.I)
BUCKETS = [("Өнөөдөр", 0, 0), ("1–3 хоног", 1, 3), ("4–7 хоног", 4, 7), ("8+ хоног", 8, 10 ** 6)]

HEAD = PatternFill("solid", fgColor="B4E5D2")
SUB1 = PatternFill("solid", fgColor="DDF3EA")
SUB2 = PatternFill("solid", fgColor="F0FAF5")
RED = PatternFill("solid", fgColor="F8D2CE")
AMBER = PatternFill("solid", fgColor="FCEBC2")
THIN = Border(bottom=Side(style="thin", color="C8D3CE"))


def norm(s):
    return re.sub(r"\s+", " ", str(s or "")).strip().lower()


def to_dt(v):
    if isinstance(v, dt.datetime):
        return v
    if isinstance(v, dt.date):
        return dt.datetime(v.year, v.month, v.day)
    if v is None or str(v).strip() == "":
        return None
    s = str(v).strip()
    for fmt in ("%Y-%m-%d %H:%M:%S", "%Y-%m-%d %H:%M", "%Y-%m-%d", "%m/%d/%Y %H:%M", "%m/%d/%Y", "%d.%m.%Y"):
        try:
            return dt.datetime.strptime(s, fmt)
        except ValueError:
            pass
    return None


def read_requests(path):
    ws = openpyxl.load_workbook(path, data_only=True).worksheets[0]
    rows = list(ws.iter_rows(values_only=True))
    for hi, r in enumerate(rows):
        names = [norm(c) for c in r]
        if "засварын төрөл" in names and "үйлчилгээ" in names:
            break
    else:
        sys.exit('Гарчгийн мөр олдсонгүй ("Засварын төрөл", "Үйлчилгээ")')

    def find(pred, label):
        for i, n in enumerate(names):
            if pred(n):
                return i
        sys.exit(f'"{label}" багана олдсонгүй')

    ix = {
        "type": find(lambda n: n == "засварын төрөл", "Засварын төрөл"),
        "tech": find(lambda n: n.startswith("техни"), "Техникийн №"),
        "loc": find(lambda n: n == "засварын байршил", "Засварын байршил"),
        "srv": find(lambda n: n == "үйлчилгээ", "Үйлчилгээ"),
        "desc": find(lambda n: n == "тайлбар", "Тайлбар"),
        "plan": find(lambda n: n.startswith("төлөвлөгөөт цаг"), "Төлөвлөгөөт цаг"),
        "load": find(lambda n: n.startswith("ачаатай"), "Ачаатай/ачаагүй"),
        "req": find(lambda n: n.startswith("хүсэлт гаргасан"), "Хүсэлт гаргасан огноо"),
    }
    out, skipped = [], 0
    for r in rows[hi + 1:]:
        if r[ix["tech"]] is None or r[ix["type"]] is None:
            if any(c is not None for c in r):
                skipped += 1  # totals row etc.
            continue
        req = to_dt(r[ix["req"]])
        if req is None:
            skipped += 1
            continue
        out.append({
            "type": str(r[ix["type"]]).strip() or BLANK,
            "tech": str(r[ix["tech"]]).strip(),
            "loc": str(r[ix["loc"]] or "").strip() or BLANK,
            "srv": str(r[ix["srv"]] or "").strip() or BLANK,
            "desc": str(r[ix["desc"]] or "").strip(),
            "plan": to_dt(r[ix["plan"]]),
            "load": str(r[ix["load"]] or "").strip(),
            "req": req,
        })
    return out, skipped


def header(ws, row, values, widths=None):
    for c, v in enumerate(values, 1):
        cell = ws.cell(row, c, v)
        cell.font = Font(bold=True)
        cell.fill = HEAD
        cell.alignment = Alignment(horizontal="center" if c > 1 else "left", vertical="center", wrap_text=True)
    if widths:
        for c, w in enumerate(widths, 1):
            ws.column_dimensions[get_column_letter(c)].width = w


def tree_table(ws, reqs, col_keys, col_of, title_cells, first_width=32):
    """Rows: type > location > service, with subtotals; columns from col_keys."""
    header(ws, 1, title_cells + ["Нийт"], [first_width] + [9] * (len(col_keys) + 1))
    ws.freeze_panes = "B2"
    tree = defaultdict(lambda: defaultdict(lambda: defaultdict(Counter)))
    for q in reqs:
        tree[q["type"]][q["loc"]][q["srv"]][col_of(q)] += 1

    def write(r, label, counter, level):
        ws.cell(r, 1, ("    " * level) + label)
        total = 0
        for c, k in enumerate(col_keys, 2):
            v = counter.get(k, 0)
            total += v
            if v:
                ws.cell(r, c, v)
        ws.cell(r, len(col_keys) + 2, total)
        fill = SUB1 if level == 0 else SUB2 if level == 1 else None
        for c in range(1, len(col_keys) + 3):
            cell = ws.cell(r, c)
            cell.border = THIN
            if fill:
                cell.fill = fill
            if level < 2:
                cell.font = Font(bold=True)
            if c > 1:
                cell.alignment = Alignment(horizontal="center")

    r = 2
    for t in sorted(tree):
        t_sum = Counter()
        for loc in tree[t].values():
            for srv in loc.values():
                t_sum.update(srv)
        write(r, t, t_sum, 0)
        r += 1
        for loc in sorted(tree[t]):
            l_sum = Counter()
            for srv in tree[t][loc].values():
                l_sum.update(srv)
            write(r, loc, l_sum, 1)
            r += 1
            for srv in sorted(tree[t][loc]):
                write(r, srv, tree[t][loc][srv], 2)
                r += 1
    g = Counter(col_of(q) for q in reqs)
    write(r, "Нийт", g, 0)
    for c in range(1, len(col_keys) + 3):
        ws.cell(r, c).fill = HEAD
    return r


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--input", required=True)
    ap.add_argument("--out", required=True)
    ap.add_argument("--today")
    a = ap.parse_args()

    reqs, skipped = read_requests(a.input)
    if not reqs:
        sys.exit("Засварын хүсэлт олдсонгүй")
    today = dt.date.fromisoformat(a.today) if a.today else max(q["req"] for q in reqs).date()
    for q in reqs:
        q["days"] = (today - q["req"].date()).days

    wb = openpyxl.Workbook()

    # 1. Pivot: columns = request day
    ws = wb.active
    ws.title = "Pivot"
    days = sorted({q["req"].date() for q in reqs})
    tree_table(ws, reqs, days, lambda q: q["req"].date(),
               ["Засварын төрөл / байршил / үйлчилгээ"] + [d.strftime("%d-%b") for d in days])

    # 2. Хүлээлт: columns = waiting buckets
    ws2 = wb.create_sheet("Хүлээлт")
    labels = [b[0] for b in BUCKETS]
    bucket = lambda q: next(b[0] for b in BUCKETS if b[1] <= q["days"] <= b[2])
    last = tree_table(ws2, reqs, labels, bucket, [f"Өнөөдөр: {today:%Y-%m-%d}"] + labels)
    for r in range(2, last + 1):
        c = ws2.cell(r, 1 + len(labels))  # 8+ column
        if c.value:
            c.fill = RED

    # 3. Анхаарах
    ws3 = wb.create_sheet("Анхаарах")
    cols = ["Ангилал", "Техникийн №", "Засварын төрөл", "Байршил", "Үйлчилгээ", "Хүсэлт гаргасан", "Хоног", "Тайлбар"]
    header(ws3, 1, cols, [26, 13, 16, 13, 13, 17, 7, 90])
    ws3.freeze_panes = "A2"
    flags = []
    for q in reqs:
        why = []
        if URGENT_WORDS.search(q["desc"]):
            why.append(("Яаралтай: аюулгүй байдал", RED))
        if norm(q["type"]) == "дуудлага":
            why.append(("Яаралтай: дуудлага", RED))
        if norm(q["loc"]) == "замд":
            why.append(("Яаралтай: замд", RED))
        if norm(q["load"]) == "ачаатай":
            why.append(("Ачаатай", AMBER))
        if q["days"] >= 8:
            why.append((f"Удаж буй ({q['days']} хоног)", AMBER))
        wo = WO_RE.findall(q["desc"])
        if wo:
            why.append(("WO: " + ", ".join(w.upper() for w in wo), None))
        if q["plan"] and q["plan"].date() < q["req"].date():
            why.append(("Алдаа: төлөвлөгөөт цаг < хүсэлтийн огноо", AMBER))
        if q["srv"] == BLANK:
            why.append(("Алдаа: үйлчилгээ хоосон", AMBER))
        for label, fill in why:
            flags.append((label, fill, q))
    order = lambda f: (0 if f[1] is RED else 1 if f[1] is AMBER else 2, -f[2]["days"])
    for r, (label, fill, q) in enumerate(sorted(flags, key=order), 2):
        vals = [label, q["tech"], q["type"], q["loc"], q["srv"], q["req"].strftime("%Y-%m-%d %H:%M"), q["days"], q["desc"]]
        for c, v in enumerate(vals, 1):
            cell = ws3.cell(r, c, v)
            cell.border = THIN
            cell.alignment = Alignment(vertical="top", wrap_text=(c == 8))
            if fill is not None and c == 1:
                cell.fill = fill
    ws3.auto_filter.ref = f"A1:H{max(2, len(flags) + 1)}"

    wb.save(a.out)
    json.dump({
        "requests": len(reqs), "skipped_rows": skipped, "today": str(today),
        "by_type": Counter(q["type"] for q in reqs), "by_location": Counter(q["loc"] for q in reqs),
        "by_service": Counter(q["srv"] for q in reqs), "by_day": {str(d): sum(1 for q in reqs if q["req"].date() == d) for d in days},
        "waiting": Counter(bucket(q) for q in reqs), "flags": Counter(f[0].split(":")[0].split(" (")[0] for f in flags),
    }, sys.stdout, ensure_ascii=False, indent=1)
    print()


if __name__ == "__main__":
    main()
