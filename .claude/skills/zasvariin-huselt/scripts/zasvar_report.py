#!/usr/bin/env python3
"""Repair request report from the Mine2TL export ("засвар ...rqst.xlsx").

Output workbook, laid out like the team's final report (nothing more):
  Sheet1  pivot layout: Count of Засварын төрөл, rows байршил > үйлчилгээ ((blank) last), columns request day
  Sheet2  report table at D5: location (navy / blue) > service (light), "Хүсэлт" total, "Нийт хүсэлт" row
  + the original export sheet(s), unchanged
JSON on stdout also has the e-mail draft numbers: requests, tech_count (unique Техникийн №)
and duplicates (same truck with more than one request).

Usage:
  zasvar_report.py --input "засвар 10-18rqst.xlsx" --out "Zasvaryn_huselt_10.18.xlsx"
"""
import argparse
import datetime as dt
import json
import re
import sys
from collections import Counter, defaultdict

import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

BLANK = "(хоосон)"

# Office accent 5 (4472C4) with the report's tints: -50 %, -25 %, +40 %
NAVY = PatternFill("solid", fgColor="203864")
BLUE = PatternFill("solid", fgColor="2F5597")
LIGHT = PatternFill("solid", fgColor="8FAADC")
PV = PatternFill("solid", fgColor="DDEBF7")
GRID = Border(*(Side(style="thin", color="000000"),) * 4)
PV_UNDER = Border(bottom=Side(style="thin", color="9BC2E6"))
PV_OVER = Border(top=Side(style="thin", color="9BC2E6"))
MN_ABC = "абвгдеёжзийклмноөпрстуүфхцчшщъыьэюя"


def mn_key(s):
    """Mongolian alphabetical order (Ө after О, Ү after У), as Excel sorts it."""
    return "".join(chr(0x2000 + MN_ABC.index(c)) if c in MN_ABC else c for c in str(s).lower())


def norm(s):
    return re.sub(r"\s+", " ", str(s or "")).strip().lower()


def to_dt(v):
    if isinstance(v, dt.datetime):
        return v
    if isinstance(v, dt.date):
        return dt.datetime(v.year, v.month, v.day)
    if v is None or str(v).strip() == "":
        return None
    s = str(v).strip()
    for fmt in ("%Y-%m-%d %H:%M:%S", "%Y-%m-%d %H:%M", "%Y-%m-%d", "%m/%d/%Y %H:%M", "%m/%d/%Y", "%d.%m.%Y"):
        try:
            return dt.datetime.strptime(s, fmt)
        except ValueError:
            pass
    return None


def read_requests(path):
    # The data may be on any sheet (e.g. after a pivot sheet); use the first
    # sheet that has the Mine2TL header row.
    found = None
    for ws in openpyxl.load_workbook(path, data_only=True).worksheets:
        rows = list(ws.iter_rows(values_only=True))
        for hi, r in enumerate(rows[:30]):
            names = [norm(c) for c in r]
            if "засварын төрөл" in names and "үйлчилгээ" in names and "засварын байршил" in names:
                found = (ws.title, rows, hi, names)
                break
        if found:
            break
    if not found:
        sys.exit('Гарчгийн мөр олдсонгүй ("Засварын төрөл", "Засварын байршил", "Үйлчилгээ")')
    _, rows, hi, names = found

    def find(pred, label):
        for i, n in enumerate(names):
            if pred(n):
                return i
        sys.exit(f'"{label}" багана олдсонгүй')

    ix = {
        "type": find(lambda n: n == "засварын төрөл", "Засварын төрөл"),
        "tech": find(lambda n: n.startswith("техни"), "Техникийн №"),
        "loc": find(lambda n: n == "засварын байршил", "Засварын байршил"),
        "srv": find(lambda n: n == "үйлчилгээ", "Үйлчилгээ"),
        "desc": find(lambda n: n == "тайлбар", "Тайлбар"),
        "plan": find(lambda n: n.startswith("төлөвлөгөөт цаг"), "Төлөвлөгөөт цаг"),
        "load": find(lambda n: n.startswith("ачаатай"), "Ачаатай/ачаагүй"),
        "req": find(lambda n: n.startswith("хүсэлт гаргасан"), "Хүсэлт гаргасан огноо"),
    }
    out, skipped = [], 0
    for r in rows[hi + 1:]:
        if r[ix["tech"]] is None or r[ix["type"]] is None:
            if any(c is not None for c in r):
                skipped += 1  # totals row etc.
            continue
        req = to_dt(r[ix["req"]])
        if req is None:
            skipped += 1
            continue
        out.append({
            "type": str(r[ix["type"]]).strip() or BLANK,
            "tech": str(r[ix["tech"]]).strip(),
            "loc": str(r[ix["loc"]] or "").strip() or BLANK,
            "srv": str(r[ix["srv"]] or "").strip() or BLANK,
            "desc": str(r[ix["desc"]] or "").strip(),
            "plan": to_dt(r[ix["plan"]]),
            "load": str(r[ix["load"]] or "").strip(),
            "req": req,
        })
    return out, skipped


def day_label(d):
    return d.strftime("%d-%b")


def pivot_sheet(ws, reqs, days):
    """Values laid out as the Excel PivotTable "Count of Засварын төрөл"."""
    last = len(days) + 2
    for c in range(1, last + 1):
        ws.cell(1, c).fill = PV
        ws.cell(1, c).border = PV_UNDER
        ws.cell(1, c).font = Font(bold=True)
    ws.cell(1, 1, "Count of Засварын төрөл")
    ws.cell(1, 2, "Column Labels")
    for c, v in enumerate(["Row Labels"] + [day_label(d) for d in days] + ["Grand Total"], 1):
        cell = ws.cell(2, c, v)
        cell.fill, cell.border, cell.font = PV, PV_UNDER, Font(bold=True)
        cell.alignment = Alignment(horizontal="left")
    r = 3

    def put(label, lst, level, total=False):
        nonlocal r
        cnt = Counter(q["req"].date() for q in lst)
        vals = [label] + [cnt.get(d) or None for d in days] + [len(lst)]
        for c, v in enumerate(vals, 1):
            cell = ws.cell(r, c, v)
            cell.font = Font(bold=level == 0)
            if total:
                cell.fill, cell.border = PV, PV_OVER
        ws.cell(r, 1).alignment = Alignment(horizontal="left", indent=level)
        r += 1

    for loc in sorted({q["loc"] for q in reqs}, key=mn_key):
        in_loc = [q for q in reqs if q["loc"] == loc]
        put(loc, in_loc, 0)
        srv = sorted({q["srv"] for q in in_loc if q["srv"] != BLANK}, key=mn_key)
        if any(q["srv"] == BLANK for q in in_loc):
            srv.append(BLANK)
        for s in srv:
            put("(blank)" if s == BLANK else s, [q for q in in_loc if q["srv"] == s], 1)
    put("Grand Total", reqs, 0, total=True)
    widths = [23.21875, 15.5546875] + [6.5546875] * (len(days) - 1) + [10.77734375]
    for c, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(c)].width = w


def report_table(ws, reqs, days, top=5, left=4):
    """Presentation table: location (dark) > service (light) rows, request-day
    columns, "Хүсэлт" total column and "Нийт хүсэлт" row. Blank services are
    counted in the location row but not listed."""
    cnt = defaultdict(lambda: defaultdict(Counter))
    for q in reqs:
        cnt[q["loc"]][q["srv"]][q["req"].date()] += 1
    ncol = len(days) + 2

    def put(r, vals, fill, white):
        for i, v in enumerate(vals):
            c = ws.cell(r, left + i, v)
            c.fill = fill
            c.border = GRID
            c.alignment = Alignment(horizontal="center", vertical="center")
            c.font = Font(bold=white, color="FFFFFF" if white else "000000")

    def row(label, counter):
        return [label] + [counter.get(d, 0) or None for d in days] + [sum(counter.values())]

    put(top, ["Засварын төрөл"] + [day_label(d) for d in days] + ["Хүсэлт"], NAVY, True)
    r = top + 1
    for loc in sorted(cnt, key=mn_key):
        total = Counter()
        for s in cnt[loc].values():
            total.update(s)
        services = sorted((s for s in cnt[loc] if s != BLANK), key=mn_key)
        # A location with no listed services (e.g. Замд) is styled like a service row.
        put(r, row(loc, total), BLUE if services else LIGHT, bool(services))
        r += 1
        for srv in services:
            put(r, row(srv, cnt[loc][srv]), LIGHT, False)
            r += 1
    put(r, row("Нийт хүсэлт", Counter(q["req"].date() for q in reqs)), NAVY, True)
    ws.column_dimensions[get_column_letter(left)].width = 15.88671875
    return r


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--input", required=True)
    ap.add_argument("--out", required=True)
    a = ap.parse_args()

    reqs, skipped = read_requests(a.input)
    if not reqs:
        sys.exit("Засварын хүсэлт олдсонгүй")

    # Keep the original export as-is and put the report sheets in front of it.
    wb = openpyxl.load_workbook(a.input)
    src_names = wb.sheetnames

    def free(name):
        n, k = name, 2
        while n in wb.sheetnames:
            n, k = f"{name} ({k})", k + 1
        return n

    days = sorted({q["req"].date() for q in reqs})
    pivot_sheet(wb.create_sheet(free("Sheet1"), 0), reqs, days)
    ws_rep = wb.create_sheet(free("Sheet2"), 1)
    report_table(ws_rep, reqs, days)
    for w in wb.worksheets:
        w.sheet_view.tabSelected = w is ws_rep
    wb.active = 1
    wb.save(a.out)

    by_tech = defaultdict(list)
    for q in reqs:
        by_tech[q["tech"]].append(q)
    dups = {t: [{"day": str(q["req"].date()), "type": q["type"], "loc": q["loc"], "srv": "" if q["srv"] == BLANK else q["srv"], "desc": q["desc"]} for q in l]
            for t, l in sorted(by_tech.items(), key=lambda x: mn_key(x[0])) if len(l) > 1}
    json.dump({
        "requests": len(reqs), "tech_count": len(by_tech), "skipped_rows": skipped,
        "by_location": Counter(q["loc"] for q in reqs), "by_service": Counter(q["srv"] for q in reqs),
        "by_day": {str(d): sum(1 for q in reqs if q["req"].date() == d) for d in days},
        "duplicates": dups, "source_sheets": src_names,
    }, sys.stdout, ensure_ascii=False, indent=1)
    print()


if __name__ == "__main__":
    main()
